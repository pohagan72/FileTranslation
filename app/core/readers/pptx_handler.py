"""PPTX read/translate implementation."""

from __future__ import annotations

import io
import logging
from typing import Iterable, List

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .base import DocumentHandler

logger = logging.getLogger(__name__)


class PptxHandler(DocumentHandler):
    extension = ".pptx"

    def extract_text(self, stream: io.BytesIO) -> str:
        return "\n".join(self.collect_segments(stream))

    def collect_segments(self, stream: io.BytesIO) -> List[str]:
        stream.seek(0)
        ppt = Presentation(stream)
        segments = [p.text for p in _iter_translatable_paragraphs(ppt)]
        stream.seek(0)
        return segments

    def apply_translations(self, stream: io.BytesIO, translations: List[str]) -> io.BytesIO:
        stream.seek(0)
        ppt = Presentation(stream)
        paragraphs = list(_iter_translatable_paragraphs(ppt))
        if len(paragraphs) != len(translations):
            raise RuntimeError(
                f"pptx translation count mismatch: {len(paragraphs)} segments "
                f"vs {len(translations)} translations"
            )
        for paragraph, translated in zip(paragraphs, translations):
            _replace_paragraph_text(paragraph, translated)

        output = io.BytesIO()
        ppt.save(output)
        output.seek(0)
        return output


def _walk_shapes(shapes) -> Iterable:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)


def _iter_translatable_paragraphs(ppt) -> Iterable:
    for slide in ppt.slides:
        for shape in _walk_shapes(slide.shapes):
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.text.strip():
                        yield paragraph
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            for paragraph in cell.text_frame.paragraphs:
                                if paragraph.text.strip():
                                    yield paragraph


def _replace_paragraph_text(paragraph, translated: str) -> None:
    if not translated or not translated.strip():
        return
    paragraph.clear()
    run = paragraph.add_run()
    run.text = translated
