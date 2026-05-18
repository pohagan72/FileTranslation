"""DOCX read/translate implementation."""

from __future__ import annotations

import io
import logging
from typing import Iterable, List

from docx import Document

from .base import DocumentHandler

logger = logging.getLogger(__name__)


class DocxHandler(DocumentHandler):
    extension = ".docx"

    def extract_text(self, stream: io.BytesIO) -> str:
        return "\n".join(self.collect_segments(stream))

    def collect_segments(self, stream: io.BytesIO) -> List[str]:
        stream.seek(0)
        doc = Document(stream)
        segments = [p.text for p in _iter_translatable_paragraphs(doc)]
        stream.seek(0)
        return segments

    def apply_translations(self, stream: io.BytesIO, translations: List[str]) -> io.BytesIO:
        stream.seek(0)
        doc = Document(stream)
        paragraphs = list(_iter_translatable_paragraphs(doc))
        if len(paragraphs) != len(translations):
            raise RuntimeError(
                f"docx translation count mismatch: {len(paragraphs)} segments "
                f"vs {len(translations)} translations"
            )
        for paragraph, translated in zip(paragraphs, translations):
            _replace_paragraph_text(paragraph, translated)

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output


def _iter_translatable_paragraphs(doc) -> Iterable:
    """Yield paragraphs with non-empty text from body + tables."""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        yield paragraph


def _replace_paragraph_text(paragraph, translated: str) -> None:
    if not translated or not translated.strip():
        return
    # Best-effort: copy font properties from the first existing run.
    font_style = {}
    if paragraph.runs:
        first = paragraph.runs[0]
        font_style["name"] = first.font.name
        font_style["size"] = first.font.size
        font_style["bold"] = first.bold
        font_style["italic"] = first.italic
        font_style["underline"] = first.underline

    paragraph.clear()
    run = paragraph.add_run(translated)
    if font_style.get("name") is not None:
        run.font.name = font_style["name"]
    if font_style.get("size") is not None:
        run.font.size = font_style["size"]
    if font_style.get("bold") is not None:
        run.bold = font_style["bold"]
    if font_style.get("italic") is not None:
        run.italic = font_style["italic"]
    if font_style.get("underline") is not None:
        run.underline = font_style["underline"]
