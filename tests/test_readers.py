"""Tests for the format-specific document handlers."""

from __future__ import annotations

import io

import pandas as pd
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.core.readers.docx_handler import DocxHandler
from app.core.readers.pptx_handler import PptxHandler
from app.core.readers.registry import get_handler, supported_extensions
from app.core.readers.xlsx_handler import XlsxHandler


def _docx_stream(paragraphs: list[str]) -> io.BytesIO:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _pptx_stream(texts: list[str]) -> io.BytesIO:
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[5])
    for text in texts:
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        tb.text_frame.text = text
    buf = io.BytesIO()
    ppt.save(buf)
    buf.seek(0)
    return buf


def _xlsx_stream(rows: list[list[str]], columns: list[str]) -> io.BytesIO:
    df = pd.DataFrame(rows, columns=columns)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False)
    buf.seek(0)
    return buf


def test_registry_has_three_handlers():
    assert set(supported_extensions()) == {".docx", ".pptx", ".xlsx"}


def test_registry_lookup_is_case_insensitive():
    assert isinstance(get_handler(".DOCX"), DocxHandler)


def test_docx_roundtrip_translates_each_paragraph():
    stream = _docx_stream(["Hello", "World"])
    handler = DocxHandler()

    segments = handler.collect_segments(stream)
    assert segments == ["Hello", "World"]

    translated = handler.apply_translations(stream, ["HELLO", "WORLD"])
    doc = Document(translated)
    assert [p.text for p in doc.paragraphs if p.text.strip()] == ["HELLO", "WORLD"]


def test_pptx_roundtrip_translates_each_text_box():
    stream = _pptx_stream(["Hello", "World"])
    handler = PptxHandler()

    segments = handler.collect_segments(stream)
    assert segments == ["Hello", "World"]

    translated = handler.apply_translations(stream, ["HELLO", "WORLD"])
    ppt = Presentation(translated)
    texts = []
    for slide in ppt.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.text.strip():
                        texts.append(paragraph.text)
    assert texts == ["HELLO", "WORLD"]


def test_xlsx_roundtrip_translates_each_cell():
    # Header row is preserved by pandas as column names, so only data cells
    # are translated. Order is row-major.
    stream = _xlsx_stream(
        rows=[["Hello", "World"], ["Foo", "Bar"]],
        columns=["A", "B"],
    )
    handler = XlsxHandler()

    segments = handler.collect_segments(stream)
    assert segments == ["Hello", "World", "Foo", "Bar"]

    translated_stream = handler.apply_translations(stream, ["HELLO", "WORLD", "FOO", "BAR"])
    df = pd.read_excel(translated_stream)
    flat = [str(v) for v in df.values.flatten() if pd.notna(v)]
    assert flat == ["HELLO", "WORLD", "FOO", "BAR"]


def test_apply_translations_rejects_count_mismatch():
    stream = _docx_stream(["A", "B"])
    handler = DocxHandler()
    import pytest

    with pytest.raises(RuntimeError, match="count mismatch"):
        handler.apply_translations(stream, ["only one"])
