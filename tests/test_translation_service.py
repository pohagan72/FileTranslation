"""Tests for the TranslationService orchestrator."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.core import TranslationError, TranslationService, UnsupportedFileType
from app.core.providers import TranslationProvider


def _make_docx(paragraphs: list[str]) -> io.BytesIO:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def test_translate_document_happy_path(service, storage, provider):
    stream = _make_docx(["Hello world"])

    result = service.translate_document(
        file_stream=stream,
        original_filename="sample.docx",
        extension=".docx",
        target_language="Spanish",
    )

    assert result.download_filename == "translated_sample.docx"
    assert result.download_url.startswith("https://signed.example/")
    assert provider.calls == ["Hello world"]

    # The translated blob is what's downloadable; original was deleted.
    assert any("translated/" in key for key in storage.objects)
    assert all("uploaded/" not in key for key in storage.objects)
    assert any("uploaded/" in key for key in storage.deleted)


def test_translate_document_rejects_unknown_extension(service):
    with pytest.raises(UnsupportedFileType):
        service.translate_document(
            file_stream=io.BytesIO(b"x"),
            original_filename="x.pdf",
            extension=".pdf",
            target_language="Spanish",
        )


class FlakyProvider(TranslationProvider):
    def __init__(self, fail_times: int) -> None:
        self.fail_times = fail_times
        self.attempts = 0

    def translate(self, text, target_language):
        self.attempts += 1
        if self.attempts <= self.fail_times:
            raise TranslationError("transient")
        return text.upper()


def test_translate_retries_on_transient_failures(storage):
    provider = FlakyProvider(fail_times=1)
    service = TranslationService(
        storage=storage, provider=provider, translation_threads=1, max_retries=2
    )

    service.translate_document(
        file_stream=_make_docx(["Hi"]),
        original_filename="a.docx",
        extension=".docx",
        target_language="Spanish",
    )

    assert provider.attempts == 2  # 1 failure + 1 success
    downloaded = storage.objects[next(k for k in storage.objects if "translated/" in k)]
    doc = Document(io.BytesIO(downloaded))
    assert [p.text for p in doc.paragraphs] == ["HI"]


def test_translate_keeps_original_text_when_provider_gives_up(storage):
    provider = FlakyProvider(fail_times=99)
    service = TranslationService(
        storage=storage, provider=provider, translation_threads=1, max_retries=1
    )

    service.translate_document(
        file_stream=_make_docx(["Hi"]),
        original_filename="a.docx",
        extension=".docx",
        target_language="Spanish",
    )

    downloaded = storage.objects[next(k for k in storage.objects if "translated/" in k)]
    doc = Document(io.BytesIO(downloaded))
    # Untranslated segments preserve the original text (logged in the service).
    assert [p.text for p in doc.paragraphs] == ["Hi"]
